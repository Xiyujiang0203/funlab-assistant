import io
import time
import zipfile
from pathlib import Path

import requests
from langchain_core.documents import Document

from app.config import get_settings

MINERU_EXTENSIONS = {".pdf", ".pptx", ".ppt"}


def _mineru_extract_text(file_path: Path) -> str:
    settings = get_settings()
    token = settings.mineru_api_token
    if not token:
        raise RuntimeError(f"MinerU token 未配置，无法解析 {file_path.name}")

    url = "https://mineru.net/api/v4/file-urls/batch"
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {token}",
    }
    payload = {
        "enable_formula": True,
        "enable_table": True,
        "language": "ch",
        "files": [
            {
                "name": file_path.name,
                "is_ocr": True,
                "data_id": "file_0",
            }
        ],
    }
    resp = requests.post(url, headers=headers, json=payload, timeout=60)
    resp.raise_for_status()
    result = resp.json()
    if result.get("code") != 0:
        raise RuntimeError(result.get("msg", "MinerU 上传失败"))

    batch_id = result["data"]["batch_id"]
    upload_url = result["data"]["file_urls"][0]
    with open(file_path, "rb") as f:
        up = requests.put(upload_url, data=f, timeout=120)
        up.raise_for_status()

    poll_url = f"https://mineru.net/api/v4/extract-results/batch/{batch_id}"
    for _ in range(60):
        poll = requests.get(poll_url, headers=headers, timeout=60)
        poll.raise_for_status()
        poll_json = poll.json()
        if poll_json.get("code") != 0:
            raise RuntimeError(poll_json.get("msg", "MinerU 查询失败"))

        items = poll_json["data"]["extract_result"]
        if items[0]["state"] == "failed":
            raise RuntimeError(f"MinerU 解析失败: {file_path.name}")
        if items[0]["state"] == "done":
            zip_resp = requests.get(items[0]["full_zip_url"], timeout=120)
            zip_resp.raise_for_status()
            with zipfile.ZipFile(io.BytesIO(zip_resp.content)) as zf:
                md_names = [n for n in zf.namelist() if n.endswith(".md")]
                if not md_names:
                    raise RuntimeError(f"MinerU 结果无 markdown: {file_path.name}")
                with zf.open(md_names[0]) as md_file:
                    return md_file.read().decode("utf-8", errors="ignore")
        time.sleep(5)

    raise RuntimeError(f"MinerU 解析超时: {file_path.name}")


def should_use_mineru(file_path: Path) -> bool:
    settings = get_settings()
    return file_path.suffix.lower() in MINERU_EXTENSIONS and bool(settings.mineru_api_token)


def _doc(source: str, suffix: str, text: str) -> list[Document]:
    return [Document(page_content=text, metadata={"source": source, "file_type": suffix})]


def _load_html(file_path: Path, source: str, suffix: str) -> list[Document]:
    from bs4 import BeautifulSoup

    html = file_path.read_text(encoding="utf-8", errors="ignore")
    text = BeautifulSoup(html, "lxml").get_text(separator="\n", strip=True)
    return _doc(source, suffix, text)


def _load_docx(file_path: Path, source: str, suffix: str) -> list[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(file_path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return _doc(source, suffix, text)


def _load_pptx(file_path: Path, source: str, suffix: str) -> list[Document]:
    from pptx import Presentation

    prs = Presentation(str(file_path))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
    return _doc(source, suffix, "\n".join(parts))


def load_file(file_path: Path) -> list[Document]:
    suffix = file_path.suffix.lower()
    source = str(file_path.resolve())

    if should_use_mineru(file_path):
        text = _mineru_extract_text(file_path)
        return [Document(page_content=text, metadata={"source": source, "file_type": suffix})]

    if suffix in {".txt", ".md"}:
        from langchain_community.document_loaders import TextLoader

        return TextLoader(str(file_path), encoding="utf-8").load()

    if suffix in {".html", ".htm"}:
        return _load_html(file_path, source, suffix)

    if suffix == ".docx":
        return _load_docx(file_path, source, suffix)

    if suffix == ".pdf":
        from langchain_community.document_loaders import PyPDFLoader

        return PyPDFLoader(str(file_path)).load()

    if suffix == ".pptx":
        return _load_pptx(file_path, source, suffix)

    if suffix == ".ppt":
        if get_settings().mineru_api_token:
            text = _mineru_extract_text(file_path)
            return [Document(page_content=text, metadata={"source": source, "file_type": suffix})]
        raise ValueError(f"不支持的文件类型: {file_path.name}")

    raise ValueError(f"不支持的文件类型: {file_path.name}")


def load_knowledge_dir(knowledge_dir: Path | None = None) -> list[Document]:
    settings = get_settings()
    root = knowledge_dir or settings.knowledge_path
    if not root.exists():
        return []

    supported = {".txt", ".md", ".html", ".htm", ".docx", ".pdf", ".pptx", ".ppt"}
    documents: list[Document] = []

    for file_path in sorted(root.rglob("*")):
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() not in supported:
            continue
        try:
            docs = load_file(file_path)
            for doc in docs:
                doc.metadata.setdefault("source", str(file_path.resolve()))
                doc.metadata.setdefault("file_type", file_path.suffix.lower())
            documents.extend(docs)
        except Exception as exc:
            print(f"[skip] {file_path}: {exc}")

    return documents
